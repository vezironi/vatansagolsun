"""Panelden Excel (.xlsx) ile toplu şehit ekleme: şablon, satır okuma, aktarım.

Her satır aktarım çekirdeğinden (`aktarim.kayit_aktar`) geçer: sitede zaten olan kişi yeniden
eklenmez, yalnızca boş alanları dolar; aynı dosya ikinci kez yüklenirse satırlar atlanır.
"""
import io
import zipfile
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import aktarim
from app.kaynaklar import HamKayit
from app.models import Durum, Hassasiyet, Kaynak, Kategori, Sehit
from app.text import isim_anahtari, katla

KAYNAK_KODU = "excel"
KAYNAK_ADI = "Excel ile toplu ekleme"
DOSYA_EN_COK_BAYT = 5 * 1024 * 1024
SATIR_EN_COK = 5000
KAYNAK_URL_EN_COK = 5
_HATA_GOSTERIM_SINIRI = 100

# (başlık, alan, açıklama) — başlıklar büyük/küçük harf ve Türkçe harf farkı gözetmeden eşleşir.
SUTUNLAR: tuple[tuple[str, str, str], ...] = (
    ("Ad", "ad", "Zorunlu. Ör. Mehmet Ali"),
    ("Soyad", "soyad", "Ör. YILMAZ"),
    ("Görev", "gorev", "Rütbe/görev. Ör. Piyade Uzman Çavuş, Polis Memuru"),
    ("Kategori", "kategori", "asker, jandarma, polis, öğretmen ya da diğer. Boşsa görevden tahmin edilir."),
    ("Kurum", "kurum", "Ör. Jandarma Genel Komutanlığı"),
    ("Memleket", "memleket", "Ör. Sivas (Zara)"),
    ("Şehadet tarihi", "tarih", "GG.AA.YYYY, YYYY-AA-GG, Excel tarihi ya da yalnız yıl (YYYY)"),
    ("Şehadet yeri", "sehadet_yeri", "Ör. Diyarbakır (Sur)"),
    ("Şehadet olayı", "sehadet_olayi", "Nasıl şehit düştüğü, 1–3 cümle"),
    ("Doğum yılı", "dogum_yili", "YYYY"),
    ("Baba adı", "baba_adi", ""),
    ("Dönem", "donem", "Ör. Hendek operasyonları (2015–2016)"),
    ("Fotoğraf URL", "fotograf_url", "Yalnızca http(s) adresi"),
    ("Kaynak URL", "kaynaklar", "Bilginin alındığı sayfa(lar); birden çoksa boşluk ya da satır başıyla ayırın"),
    ("Not", "not_", "Panelde kaynağın altında görünür; sitede gösterilmez"),
)
_BASLIKLAR = {katla(baslik): alan for baslik, alan, _ in SUTUNLAR}
_EK_BASLIKLAR = {"guven": "guven", "rutbe": "gorev", "rutbe/gorev": "gorev", "isim": "ad", "soyadi": "soyad"}

_KATEGORILER = {
    "asker": Kategori.asker, "tsk": Kategori.asker, "jandarma": Kategori.jandarma, "polis": Kategori.polis,
    "emniyet": Kategori.polis, "ogretmen": Kategori.ogretmen, "diger": Kategori.diger, "korucu": Kategori.diger,
}
_POLIS_GOREVLERI = ("polis", "komiser", "emniyet", "bekci")
_URL = re.compile(r"https?://\S+")
_IZLENEN = ("ad", "soyad", "gorev", "kategori", "kurum", "memleket", "baba_adi", "dogum_yili", "donem",
            "sehadet_yeri", "sehadet_olayi", "fotograf_url", "sehadet_tarihi", "yil")


class DosyaHatasi(ValueError):
    """Dosya hiç okunamadı (biçim, boyut, başlık)."""


@dataclass
class Satir:
    no: int  # Excel satır numarası
    ham: HamKayit | None = None
    hata: str | None = None
    ek_kaynaklar: list[str] = field(default_factory=list)


@dataclass
class Ozet:
    eklenen: int = 0
    birlesen: int = 0
    tamamlanan: int = 0
    atlanan: int = 0
    hatali: int = 0
    hatalar: list[str] = field(default_factory=list)
    yeni_kayitlar: list[int] = field(default_factory=list)

    @property
    def toplam(self) -> int:
        return self.eklenen + self.birlesen + self.tamamlanan + self.atlanan + self.hatali


def sablon() -> bytes:
    """Boş şablon: ilk sayfada başlıklar, ikinci sayfada sütun açıklamaları."""
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Şehitler"
    sayfa.append([baslik for baslik, _, _ in SUTUNLAR])
    for hucre in sayfa[1]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="B71C1C")
        hucre.alignment = Alignment(vertical="center")
    for i, (baslik, alan, _) in enumerate(SUTUNLAR, start=1):
        genislik = 60 if alan in ("sehadet_olayi", "kaynaklar") else max(14, len(baslik) + 4)
        sayfa.column_dimensions[sayfa.cell(row=1, column=i).column_letter].width = genislik
    sayfa.freeze_panes = "A2"
    aciklama = kitap.create_sheet("Açıklama")
    aciklama.append(["Sütun", "Açıklama"])
    for baslik, _, metin in SUTUNLAR:
        aciklama.append([baslik, metin])
    aciklama.append([])
    aciklama.append(["", "Yalnızca ilk sayfa okunur. İlk satır başlıktır; sütunların sırası önemli değildir."])
    aciklama.append(["", "Sitede zaten olan kişi yeniden eklenmez: yalnızca boş bilgileri doldurulur."])
    aciklama.column_dimensions["A"].width = 18
    aciklama.column_dimensions["B"].width = 90
    for hucre in aciklama[1]:
        hucre.font = Font(bold=True)
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


def _metin(deger) -> str | None:
    if deger is None:
        return None
    if isinstance(deger, float) and deger.is_integer():
        deger = int(deger)
    metin = str(deger).strip()
    return metin or None


def _tarih(deger) -> tuple[date | None, int | None, Hassasiyet]:
    if deger is None or deger == "":
        return None, None, Hassasiyet.bilinmiyor
    if isinstance(deger, datetime):
        deger = deger.date()
    if isinstance(deger, date):
        return deger, deger.year, Hassasiyet.tam
    metin = _metin(deger) or ""
    if re.fullmatch(r"\d{4}", metin):
        return None, int(metin), Hassasiyet.yil
    for kalip, sira in ((r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", "gay"), (r"(\d{4})-(\d{1,2})-(\d{1,2})", "yag")):
        if m := re.fullmatch(kalip, metin):
            parca = dict(zip(sira, map(int, m.groups())))
            try:
                gun = date(parca["y"], parca["a"], parca["g"])
            except ValueError:
                break
            return gun, gun.year, Hassasiyet.tam
    raise ValueError(f"Şehadet tarihi anlaşılamadı: {metin!r} (GG.AA.YYYY ya da YYYY yazın)")


def _yil(deger, ad: str) -> int | None:
    metin = _metin(deger)
    if metin is None:
        return None
    if not re.fullmatch(r"\d{4}", metin):
        raise ValueError(f"{ad} dört haneli bir yıl olmalı: {metin!r}")
    return int(metin)


def _kategori(deger, gorev: str | None, kurum: str | None) -> Kategori:
    metin = katla(_metin(deger) or "")
    if metin:
        if metin not in _KATEGORILER:
            raise ValueError(f"Kategori anlaşılamadı: {deger!r} (asker, jandarma, polis, öğretmen, diğer)")
        return _KATEGORILER[metin]
    g, k = katla(gorev or ""), katla(kurum or "")
    if any(p in g for p in _POLIS_GOREVLERI) or "emniyet" in k:
        return Kategori.polis
    if "jandarma" in k or "jandarma" in g:
        return Kategori.jandarma
    if "korucu" in g or "korucu" in k:
        return Kategori.diger
    return Kategori.asker


def _dis_kimlik(ad: str, soyad: str | None, gorev: str | None) -> str:
    return f"{isim_anahtari(ad, soyad, None)}|{katla(gorev or '')}"


def _satir_oku(no: int, degerler: dict[str, object]) -> Satir:
    ad = _metin(degerler.get("ad"))
    if not ad or not any(c.isalpha() for c in ad):
        return Satir(no, hata="Ad boş")
    soyad, gorev, kurum = (_metin(degerler.get(a)) for a in ("soyad", "gorev", "kurum"))
    try:
        tarih, yil, hassasiyet = _tarih(degerler.get("tarih"))
        dogum_yili = _yil(degerler.get("dogum_yili"), "Doğum yılı")
        kategori = _kategori(degerler.get("kategori"), gorev, kurum)
    except ValueError as hata:
        return Satir(no, hata=str(hata))
    urller = _URL.findall(_metin(degerler.get("kaynaklar")) or "")[:KAYNAK_URL_EN_COK]
    notlar = [f"Not: {n}" for n in [_metin(degerler.get("not_"))] if n]
    if guven := _metin(degerler.get("guven")):
        notlar.insert(0, f"Güven: {guven}")
    ham = HamKayit(
        kaynak_kodu=KAYNAK_KODU, dis_kimlik=_dis_kimlik(ad, soyad, gorev), ad=ad, soyad=soyad,
        tarih=tarih, hassasiyet=hassasiyet, yil=yil, gorev=gorev, kategori=kategori, kurum=kurum,
        memleket=_metin(degerler.get("memleket")), baba_adi=_metin(degerler.get("baba_adi")),
        dogum_yili=dogum_yili, donem=_metin(degerler.get("donem")),
        sehadet_yeri=_metin(degerler.get("sehadet_yeri")), sehadet_olayi=_metin(degerler.get("sehadet_olayi")),
        fotograf_url=_metin(degerler.get("fotograf_url")),
        kaynak_adi=KAYNAK_ADI, kaynak_url=urller[0] if urller else None,
        metin=" · ".join(notlar) or None, resmi_teyit=False,
    )
    return Satir(no, ham=ham, ek_kaynaklar=urller[1:])


def oku(veri: bytes) -> list[Satir]:
    """İlk sayfayı okur. Dosya bütünüyle okunamazsa DosyaHatasi; tek satır hatası Satir.hata'ya yazılır."""
    if len(veri) > DOSYA_EN_COK_BAYT:
        raise DosyaHatasi(f"Dosya çok büyük (en çok {DOSYA_EN_COK_BAYT // (1024 * 1024)} MB).")
    try:
        kitap = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as hata:
        raise DosyaHatasi("Dosya okunamadı: .xlsx biçiminde bir Excel dosyası yükleyin.") from hata
    try:
        satirlar = kitap.worksheets[0].iter_rows(values_only=True)
        basliklar = next(satirlar, None) or ()
        alanlar = [_BASLIKLAR.get(katla(str(b or ""))) or _EK_BASLIKLAR.get(katla(str(b or ""))) for b in basliklar]
        if "ad" not in alanlar:
            raise DosyaHatasi("İlk satırda \"Ad\" başlığı bulunamadı. Şablonu indirip onu doldurun.")
        sonuc: list[Satir] = []
        for no, degerler in enumerate(satirlar, start=2):
            if all(d is None or str(d).strip() == "" for d in degerler):
                continue
            if len(sonuc) >= SATIR_EN_COK:
                raise DosyaHatasi(f"Çok fazla satır (en çok {SATIR_EN_COK}). Dosyayı bölüp yükleyin.")
            sozluk = {a: d for a, d in zip(alanlar, degerler) if a}
            sonuc.append(_satir_oku(no, sozluk))
        return sonuc
    finally:
        kitap.close()


def aktar(s: Session, satirlar: list[Satir], *, durum: Durum = Durum.onay_bekliyor) -> Ozet:
    """Satırları kaydeder ve commit eder. Yeni eklenen kayıtların id'leri Ozet.yeni_kayitlar'dadır.

    Daha önce yüklenmiş satır yeniden eklenmez; kayıtta boş kalan bilgiler doldurulur ve yeni kaynak
    adresleri eklenir (önce isimlerle yüklenen liste, araştırıldıkça yeniden yüklenebilir).
    """
    ozet = Ozet()

    def hata(no: int, mesaj: str) -> None:
        ozet.hatali += 1
        if len(ozet.hatalar) < _HATA_GOSTERIM_SINIRI:
            ozet.hatalar.append(f"Satır {no}: {mesaj}")

    for satir in satirlar:
        if satir.ham is None:
            hata(satir.no, satir.hata or "okunamadı")
            continue
        onceki = s.scalar(select(Kaynak).where(
            Kaynak.kaynak_kodu == KAYNAK_KODU,
            Kaynak.dis_kimlik == aktarim.dis_kimlik_sinirla(satir.ham.dis_kimlik),
        ))
        if onceki is not None:
            if _tamamla(s, onceki, satir):
                ozet.tamamlanan += 1
            else:
                ozet.atlanan += 1
            continue
        sonuc = aktarim.kayit_aktar(s, satir.ham, yeni_durum=durum, ayni_kisi_olabilir=_ayni_kisi_olabilir)
        if sonuc.durum == "hatali":
            hata(satir.no, sonuc.mesaj or "kaydedilemedi")
            continue
        if sonuc.durum == "atlandi":  # aynı dosyada aynı kişi iki kez
            ozet.atlanan += 1
            continue
        if sonuc.durum == "yeni":
            ozet.eklenen += 1
            ozet.yeni_kayitlar.append(sonuc.sehit_id)
        else:
            ozet.birlesen += 1
        _kaynaklari_ekle(s, s.get(Sehit, sonuc.sehit_id), satir, satir.ek_kaynaklar)
    s.commit()
    return ozet


def _ayni_kisi_olabilir(mevcut: Sehit, ham: HamKayit) -> bool:
    """Aynı adlı kayıtla birleştirmeden önce: polis/asker ayrımı ya da birbirini içermeyen rütbeler
    (ör. "Astsubay" ile "Uzman Çavuş") farklı kişi demektir. "Uzman Çavuş" ile "Piyade Uzman Çavuş"
    çelişmez. Şüphede yeni kayıt açılır; panelde iki kayıt birleştirilebilir."""
    kategoriler = (Kategori(mevcut.kategori), Kategori(ham.kategori))
    if Kategori.diger not in kategoriler and (kategoriler[0] == Kategori.polis) != (kategoriler[1] == Kategori.polis):
        return False
    a, b = katla(mevcut.gorev or ""), katla(ham.gorev or "")
    return not (a and b and a not in b and b not in a)


def _tamamla(s: Session, kaynak: Kaynak, satir: Satir) -> bool:
    """Daha önce yüklenmiş satırın kaydında boş alanları doldurur; bir şey değiştiyse True."""
    sehit = kaynak.sehit
    try:
        with s.begin_nested():
            once = {a: getattr(sehit, a) for a in (*_IZLENEN, "tarih_hassasiyeti")}
            aktarim._birlestir(sehit, aktarim.ham_to_veri(satir.ham))  # öncelik eşit: yalnızca boşlar dolar
            if satir.ham.metin and not kaynak.metin:
                kaynak.metin = satir.ham.metin
            urller = [u for u in [satir.ham.kaynak_url, *satir.ek_kaynaklar] if u]
            if urller and not kaynak.url:
                kaynak.url = urller.pop(0)
            bilinen = {k.url for k in sehit.kaynaklar}
            yeni_urller = [u for u in urller if u not in bilinen]
            _kaynaklari_ekle(s, sehit, satir, yeni_urller)
            degisti = bool(yeni_urller) or once != {a: getattr(sehit, a) for a in once} or s.is_modified(kaynak)
            s.flush()
    except Exception:  # noqa: BLE001 - tek satır tüm yüklemeyi düşürmemeli
        return False
    return degisti


def _kaynaklari_ekle(s: Session, sehit: Sehit, satir: Satir, urller: list[str]) -> None:
    for url in urller:
        s.add(Kaynak(
            sehit=sehit, url=url, kaynak_adi=KAYNAK_ADI, kaynak_kodu=KAYNAK_KODU,
            dis_kimlik=aktarim.dis_kimlik_sinirla(f"{satir.ham.dis_kimlik}#{url}"),
        ))
    s.flush()
