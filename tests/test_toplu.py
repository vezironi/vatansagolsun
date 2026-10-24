import io
import json
from datetime import date, datetime

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker

from app import toplu
from app.dates import bugun
from app.models import Durum, Hassasiyet, Kategori, Sehit, Webhook, WebhookDurumu, WebhookTuru

BASLIKLAR = ["Ad", "Soyad", "Görev", "Kategori", "Kurum", "Memleket", "Şehadet tarihi", "Şehadet yeri",
             "Şehadet olayı", "Kaynak URL", "Güven"]


def _xlsx(satirlar, basliklar=BASLIKLAR) -> bytes:
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.append(basliklar)
    for satir in satirlar:
        sayfa.append(satir)
    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()


def _kayitlar(session) -> list[Sehit]:
    return list(session.scalars(select(Sehit).order_by(Sehit.id)))


ORNEK = [
    ["Bünyamin", "TORĞUT", "Polis Memuru", "", "Emniyet Genel Müdürlüğü", "Elazığ", "2015", "Siirt",
     "Eruh yolunda ateş açıldı.", "https://a.example/1 https://b.example/2", "yuksek"],
    ["Mehmet Ali", "SARAK", "Uzman Çavuş", "jandarma", "", "", datetime(2016, 1, 5), "Diyarbakır (Sur)",
     None, None, None],
    ["Yakup", "AKTÜRK", "Korucu", "", "", "", "21.08.2015", None, None, None, "bulunamadi"],
]


# --- okuma ------------------------------------------------------------------------------------

def test_satirlar_okunur_tarih_ve_kategori_cozulur():
    satirlar = toplu.oku(_xlsx(ORNEK))
    assert [s.hata for s in satirlar] == [None, None, None]
    b, m, y = (s.ham for s in satirlar)
    assert (b.hassasiyet, b.yil, b.tarih) == (Hassasiyet.yil, 2015, None)
    assert b.kategori == Kategori.polis  # boş kategori görevden tahmin edilir
    assert b.kaynak_url == "https://a.example/1" and satirlar[0].ek_kaynaklar == ["https://b.example/2"]
    assert b.metin == "Güven: yuksek" and b.resmi_teyit is False and b.kaynak_kodu == "excel"
    assert (m.hassasiyet, m.tarih, m.kategori) == (Hassasiyet.tam, date(2016, 1, 5), Kategori.jandarma)
    assert (y.tarih, y.kategori) == (date(2015, 8, 21), Kategori.diger)


def test_basliklar_harf_ve_sira_farkini_tolere_eder():
    satirlar = toplu.oku(_xlsx([["2016-03-02", "ali", "Er"]], ["sehadet tarihi", "AD", "rütbe"]))
    assert satirlar[0].ham.ad == "ali" and satirlar[0].ham.gorev == "Er"
    assert satirlar[0].ham.tarih == date(2016, 3, 2)


@pytest.mark.parametrize("satir, parca", [
    (["", "X"], "Ad boş"),
    (["Ali", "X", "", "", "", "", "geçen yıl"], "Şehadet tarihi anlaşılamadı"),
    (["Ali", "X", "", "", "", "", "31.02.2016"], "Şehadet tarihi anlaşılamadı"),
    (["Ali", "X", "", "asteğmen"], "Kategori anlaşılamadı"),
])
def test_hatali_satir_tum_dosyayi_dusurmez(satir, parca):
    satirlar = toplu.oku(_xlsx([satir, ["Veli", "Y"]]))
    assert parca in satirlar[0].hata and satirlar[0].no == 2
    assert satirlar[1].ham.ad == "Veli"


def test_bos_satirlar_atlanir_ve_satir_numarasi_korunur():
    satirlar = toplu.oku(_xlsx([["Ali", "X"], [None, None], ["", "Y"]]))
    assert [s.no for s in satirlar] == [2, 4]


@pytest.mark.parametrize("veri, parca", [
    (b"duz metin", ".xlsx"),
    (b"x" * (toplu.DOSYA_EN_COK_BAYT + 1), "çok büyük"),
], ids=["xlsx-degil", "cok-buyuk"])
def test_okunamayan_dosya(veri, parca):
    with pytest.raises(toplu.DosyaHatasi, match=parca):
        toplu.oku(veri)


def test_ad_basligi_yoksa_hata():
    with pytest.raises(toplu.DosyaHatasi, match="Ad"):
        toplu.oku(_xlsx([["x"]], ["İsim soyisim"]))


def test_sablon_basliklari_okuyucuyla_uyumlu():
    sayfa = load_workbook(io.BytesIO(toplu.sablon())).worksheets[0]
    assert [h.value for h in sayfa[1]] == [b for b, _, _ in toplu.SUTUNLAR]
    assert toplu.oku(toplu.sablon()) == []


# --- aktarım ------------------------------------------------------------------------------------

def test_aktarim_onay_bekler_kaynak_ekler_ve_ikinci_yukleme_atlanir(session):
    ozet = toplu.aktar(session, toplu.oku(_xlsx(ORNEK)))
    assert (ozet.eklenen, ozet.birlesen, ozet.atlanan, ozet.hatali) == (3, 0, 0, 0)
    kayitlar = _kayitlar(session)
    assert {k.durum for k in kayitlar} == {Durum.onay_bekliyor}
    b = kayitlar[0]
    assert (b.sehadet_yeri, b.sehadet_olayi, b.memleket) == ("Siirt", "Eruh yolunda ateş açıldı.", "Elazığ")
    assert sorted(k.url for k in b.kaynaklar) == ["https://a.example/1", "https://b.example/2"]
    ozet = toplu.aktar(session, toplu.oku(_xlsx(ORNEK)))
    assert (ozet.eklenen, ozet.atlanan) == (0, 3)
    assert len(_kayitlar(session)) == 3


def test_yeniden_yukleme_bos_alanlari_doldurur_doluyu_ezmez(session):
    toplu.aktar(session, toplu.oku(_xlsx([["Yakup", "AKTÜRK", "Korucu", "", "", "Van"]])))
    zengin = ["Yakup", "AKTÜRK", "Korucu", "", "Köy korucusu", "Hakkâri", "2015-08-21", "Van (Başkale)",
              "Pusuda şehit oldu.", "https://a.example/1 https://b.example/2", "orta"]
    ozet = toplu.aktar(session, toplu.oku(_xlsx([zengin])))
    assert (ozet.eklenen, ozet.tamamlanan, ozet.atlanan) == (0, 1, 0)
    (k,) = _kayitlar(session)
    assert (k.memleket, k.kurum, k.sehadet_tarihi, k.sehadet_yeri) == ("Van", "Köy korucusu", date(2015, 8, 21), "Van (Başkale)")
    assert sorted(x.url for x in k.kaynaklar) == ["https://a.example/1", "https://b.example/2"]
    assert k.kaynaklar[0].metin == "Güven: orta"
    ozet = toplu.aktar(session, toplu.oku(_xlsx([zengin])))
    assert (ozet.tamamlanan, ozet.atlanan) == (0, 1)
    assert len(k.kaynaklar) == 2


def test_sitede_olan_kisi_yeniden_eklenmez_bos_alanlari_dolar(session, sehit_ekle):
    mevcut = sehit_ekle(ad="Mehmet Ali", soyad="SARAK", sehadet_tarihi=date(2016, 1, 5), kaynak_kodu="egm",
                        gorev="Jandarma Uzman Çavuş")
    ozet = toplu.aktar(session, toplu.oku(_xlsx(ORNEK[1:2])))
    assert (ozet.eklenen, ozet.birlesen) == (0, 1)
    session.refresh(mevcut)
    assert mevcut.sehadet_yeri == "Diyarbakır (Sur)"
    assert mevcut.gorev == "Jandarma Uzman Çavuş"  # resmî kaynağın bilgisi ezilmez
    assert mevcut.durum == Durum.yayinda


def test_ayni_adli_farkli_rutbeler_birlesmez(session, sehit_ekle):
    sehit_ekle(ad="Uğur", soyad="YILDIZ", gorev="Polis Memuru", kategori="polis", sehadet_tarihi=None,
               tarih_hassasiyeti="bilinmiyor", donem="Hendek")
    satirlar = [["Uğur", "YILDIZ", "Er", "", "", "", "", "", "", "", ""],
                ["Murat", "DEMİRCİ", "Astsubay"], ["Murat", "DEMİRCİ", "Uzman Çavuş"],
                ["Ali", "KAYA", "Uzman Çavuş"], ["Ali", "KAYA", "Piyade Uzman Çavuş"]]
    ozet = toplu.aktar(session, toplu.oku(_xlsx(satirlar, BASLIKLAR[:3] + BASLIKLAR[3:])))
    assert (ozet.eklenen, ozet.birlesen) == (4, 1)  # yalnızca rütbesi birbirini içeren Ali KAYA birleşir
    assert len(_kayitlar(session)) == 5


def test_hatali_satirlar_ozette(session):
    ozet = toplu.aktar(session, toplu.oku(_xlsx([["", "X"], ["Ali", "Y"]])))
    assert (ozet.eklenen, ozet.hatali) == (1, 1)
    assert ozet.hatalar == ["Satır 2: Ad boş"]


# --- panel ----------------------------------------------------------------------------------------

def _yukle(panel_client, csrf, veri, **form):
    return panel_client.post(
        "/panel/toplu", data={"csrf": csrf(panel_client), **form},
        files={"dosya": ("liste.xlsx", veri, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_panel_sayfasi_ve_sablon(panel_client):
    h = panel_client.get("/panel/toplu").text
    assert "Toplu şehit ekle" in h and 'href="/panel/toplu"' in h
    r = panel_client.get("/panel/toplu/sablon.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).worksheets[0]["A1"].value == "Ad"


def test_panel_toplu_yukleme(panel_client, csrf, session):
    r = _yukle(panel_client, csrf, _xlsx(ORNEK))
    assert r.status_code == 200 and "3 yeni kayıt eklendi" in r.text and "onay bekliyor" in r.text
    assert len(_kayitlar(session)) == 3


def test_panel_okunamayan_dosyada_hata(panel_client, csrf, session):
    r = _yukle(panel_client, csrf, b"duz metin")
    assert r.status_code == 400 and ".xlsx" in r.text
    assert _kayitlar(session) == []


def test_panel_toplu_giris_ve_csrf_ister(client, panel_client, session):
    assert client.get("/panel/toplu", follow_redirects=False).status_code == 303
    assert client.get("/panel/toplu/sablon.xlsx", follow_redirects=False).status_code == 303
    r = panel_client.post("/panel/toplu", data={"csrf": "yanlis"}, files={"dosya": ("a.xlsx", _xlsx(ORNEK))})
    assert r.status_code == 403 and _kayitlar(session) == []


def test_panel_dogrudan_yayina_yeni_sehiti_duyurur(app, panel_client, csrf, session, monkeypatch):
    from app import webhooks

    giden = []
    monkeypatch.setattr(webhooks, "gonder", lambda url, govde, basliklar: giden.append(govde) or httpx.Response(204))
    app.state.session_factory = sessionmaker(session.get_bind())
    session.add(Webhook(url="https://discord.com/api/webhooks/1/abc", tur=WebhookTuru.discord,
                        durum=WebhookDurumu.aktif, silme_anahtari_ozeti="x" * 64, imza_anahtari="y" * 64))
    session.commit()
    r = _yukle(panel_client, csrf, _xlsx([["Ali", "VELİ", "Er", "", "", "", bugun()], ORNEK[0]]), durum="yayinda")
    assert "yayına alındı" in r.text
    assert {k.durum for k in _kayitlar(session)} == {Durum.yayinda}
    assert len(giden) == 1  # yalnızca son 7 günün şehidi duyurulur
    assert "VELİ" in json.loads(giden[0])["embeds"][0]["title"]


# --- toplu onay -----------------------------------------------------------------------------------

def test_panel_secilenleri_onaylar_ve_reddeder(panel_client, csrf, session, sehit_ekle):
    a, b, c = (sehit_ekle(ad=ad, durum="onay_bekliyor") for ad in ("Ali", "Veli", "Can"))
    yayinda = sehit_ekle(ad="Eski")
    h = panel_client.get("/panel").text
    assert 'action="/panel/toplu-islem"' in h and f'name="id" value="{a.id}"' in h
    r = panel_client.post("/panel/toplu-islem", data={"csrf": csrf(panel_client), "islem": "onayla",
                                                      "id": [a.id, b.id, yayinda.id, 999999]})
    assert "2 kayıt yayına alındı." in r.text
    panel_client.post("/panel/toplu-islem", data={"csrf": csrf(panel_client), "islem": "reddet", "id": [c.id]})
    for k in (a, b, c, yayinda):
        session.refresh(k)
    assert (a.durum, b.durum, c.durum, yayinda.durum) == (Durum.yayinda, Durum.yayinda, Durum.reddedildi, Durum.yayinda)


def test_panel_toplu_islem_guveni_isaretler(panel_client, csrf, session):
    toplu.aktar(session, toplu.oku(_xlsx(ORNEK)))
    h = panel_client.get("/panel").text
    assert 'data-guven="yuksek"' in h and 'data-guven="bulunamadi"' in h


def test_panel_toplu_islem_csrf_ve_gecersiz_islem(panel_client, csrf, session, sehit_ekle):
    a = sehit_ekle(durum="onay_bekliyor")
    assert panel_client.post("/panel/toplu-islem", data={"csrf": "x", "islem": "onayla", "id": [a.id]}).status_code == 403
    assert panel_client.post("/panel/toplu-islem", data={"csrf": csrf(panel_client), "islem": "sil",
                                                         "id": [a.id]}).status_code == 400
    r = panel_client.post("/panel/toplu-islem", data={"csrf": csrf(panel_client), "islem": "onayla"})
    assert "0 kayıt yayına alındı." in r.text
    session.refresh(a)
    assert a.durum == Durum.onay_bekliyor
